"""Nur oqiw mathematics (grades 2–7) API.

Separate from Telegram bot and legacy school group data.
All updates are transactional and bump a monotonic revision. Monthly ranking
is dense: 1,1,2,3 rather than competition ranking 1,1,3,4.
"""
import calendar
import csv
import datetime as dt
import io
import os
import re
import time
import unicodedata
import zipfile
from collections import defaultdict
from zoneinfo import ZoneInfo

from flask import Response, jsonify, request, stream_with_context
from openpyxl import load_workbook

TZ = ZoneInfo("Asia/Tashkent")
NAME_LIMIT = 120
VALID_STATUS = {"present", "dq"}
MAX_ROWS = 450
REQUIRED_HEADERS = {
    "ism": "name", "familiya": "name", "fio": "name",
    "ismfamiliya": "name", "fullname": "name", "oquvchi": "name",
    "student": "name", "name": "name",
    "ball": "correct", "score": "correct", "togri": "correct",
    "togrisoni": "correct", "correct": "correct",
    "correctanswers": "correct", "natija": "correct",
    "dq": "status", "holat": "status", "status": "status",
    "davomat": "status", "attendance": "status",
    "sinf": "grade", "grade": "grade", "class": "grade",
}

def register_math(app, db, admin_only, bad):
    def date_value(value):
        try:
            return dt.date.fromisoformat(str(value))
        except (TypeError, ValueError):
            raise ValueError("Sana YYYY-MM-DD formatida kiritilsin")

    def month_value(value):
        if not re.fullmatch(r"\d{4}-\d{2}", str(value or "")):
            raise ValueError("Oy YYYY-MM bo‘lishi kerak")
        return date_value(str(value) + "-01")

    def this_month():
        now = dt.datetime.now(TZ).date()
        return dt.date(now.year, now.month, 1)

    def text(obj, field, maximum=120):
        value = obj.get(field)
        if not isinstance(value, str) or not (1 <= len(value.strip()) <= maximum):
            raise ValueError(field + " kiritilishi kerak")
        return value.strip()

    def bump(conn):
        conn.execute("""UPDATE math_revision SET revision=revision+1,
                        updated_at=now() WHERE singleton=true""")
        # Transactional event. A future dedicated VPS can LISTEN and fan out.
        conn.execute("SELECT pg_notify('nur_math_changed','updated')")

    def revision(conn):
        row = conn.execute("SELECT revision FROM math_revision WHERE singleton=true").fetchone()
        return int(row["revision"])

    def normalize(v):
        v = unicodedata.normalize("NFKD", str(v or "")).casefold()
        return "".join(c for c in v if c.isalnum() and not unicodedata.combining(c))

    def parse_row(raw, question_count):
        if not isinstance(raw, dict):
            raise ValueError("Natija qatori obyekt bo‘lishi kerak")
        name = text(raw, "name", NAME_LIMIT)
        status = str(raw.get("status") or "present").strip().lower()
        if status in ("absent", "kelmadi", "qatnashmadi", "d", "dq"):
            status = "dq"
        if status not in VALID_STATUS:
            raise ValueError(name + ": holat present yoki dq bo‘lishi kerak")
        grade = raw.get("grade")
        if grade in ("", None):
            grade = None
        else:
            try: grade = int(grade)
            except (TypeError, ValueError): raise ValueError(name + ": sinf raqami noto‘g‘ri")
            if not 2 <= grade <= 7:
                raise ValueError(name + ": 2–7 sinf oralig‘i")
        correct = raw.get("correct")
        if status == "dq":
            correct = None
        else:
            if correct in ("", None):
                raise ValueError(name + ": ball noaniq. DQ, 0 yoki haqiqiy ballni tanlang")
            try:
                correct = int(correct)
            except (TypeError, ValueError):
                raise ValueError(name + ": to‘g‘ri javob butun son bo‘lishi kerak")
            if not 0 <= correct <= question_count:
                raise ValueError(name + ": ball 0–" + str(question_count) + " oralig‘ida")
        return {"name":name, "grade":grade, "status":status, "correct":correct}

    def leaderboard(conn, month):
        month_end = dt.date(month.year + (month.month == 12),
                            month.month % 12 + 1, 1)
        rows = conn.execute("""SELECT st.id,st.full_name,st.grade,
             COALESCE(SUM(r.correct_answers) FILTER (
                  WHERE t.held_on>=%s AND t.held_on<%s AND r.status='present'),0)::int AS monthly_points,
             COALESCE(SUM(r.correct_answers) FILTER (WHERE r.status='present'),0)::int AS total_points,
             count(r.test_id) FILTER (WHERE t.held_on>=%s AND t.held_on<%s)::int AS monthly_tests,
             count(r.test_id) FILTER (WHERE t.held_on>=%s AND t.held_on<%s AND r.status='present')::int AS attended,
             count(r.test_id) FILTER (WHERE t.held_on>=%s AND t.held_on<%s AND r.status='dq')::int AS dq,
             count(r.test_id) FILTER (WHERE r.status='present')::int AS all_tests
             FROM math_students st
             LEFT JOIN math_results r ON r.student_id=st.id
             LEFT JOIN math_tests t ON t.id=r.test_id
             WHERE NOT st.archived
             GROUP BY st.id ORDER BY monthly_points DESC,st.full_name,st.id""",
             (month,month_end,month,month_end,month,month_end,month,month_end)).fetchall()
        payments = conn.execute("""SELECT student_id, paid, amount_uzs, paid_at
            FROM math_monthly_payments WHERE month=%s""", (month,)).fetchall()
        payment_map={row["student_id"]:row for row in payments}
        previous=None
        rank=0
        out=[]
        for src in rows:
            s=dict(src)
            if s["monthly_tests"]>0:
                if previous != s["monthly_points"]:
                    rank+=1
                    previous=s["monthly_points"]
                s["rank"]=rank
            else:
                s["rank"]=None
            s["attendance_total"]=s["monthly_tests"]
            p=payment_map.get(s["id"])
            s["paid"]=bool(p and p["paid"])
            s["amount_uzs"]=p["amount_uzs"] if p else None
            out.append(s)
        return out

    def snapshot(conn, day, month):
        students=leaderboard(conn,month)
        dates=conn.execute("""SELECT t.id,t.held_on,t.title,t.question_count,
            count(r.student_id)::int AS entered
            FROM math_tests t LEFT JOIN math_results r ON r.test_id=t.id
            GROUP BY t.id ORDER BY t.held_on DESC,t.id DESC LIMIT 90""").fetchall()
        selected=[x for x in dates if x["held_on"]==day]
        tests=[]
        for test in selected:
            marks=conn.execute("""SELECT st.id AS student_id,st.full_name,st.grade,
              r.status,r.correct_answers
              FROM math_results r JOIN math_students st ON st.id=r.student_id
              WHERE r.test_id=%s AND NOT st.archived
              ORDER BY r.status DESC,r.correct_answers DESC NULLS LAST,st.full_name""",
              (test["id"],)).fetchall()
            tests.append({"id":test["id"],"title":test["title"],
                          "questions":test["question_count"],"rows":[dict(x) for x in marks]})
        return {"revision":revision(conn),"day":day.isoformat(),"month":month.strftime("%Y-%m"),
                "dates":[{"date":x["held_on"].isoformat(),"title":x["title"],
                          "questions":x["question_count"],"entered":x["entered"]} for x in dates],
                "tests":tests,"students":students,"total_students":len(students)}

    @app.get("/api/math/overview")
    def math_overview():
        day=date_value(request.args.get("date") or dt.datetime.now(TZ).date().isoformat())
        month=month_value(request.args.get("month") or day.strftime("%Y-%m"))
        with db() as conn:
            result=snapshot(conn,day,month)
        return jsonify(result)

    @app.get("/api/math/student/<int:student_id>")
    def math_profile(student_id):
        with db() as conn:
            s=conn.execute("""SELECT id,full_name,grade FROM math_students
                 WHERE id=%s AND NOT archived""", (student_id,)).fetchone()
            if not s: return bad("O‘quvchi topilmadi",404)
            history=conn.execute("""SELECT t.held_on,t.title,t.question_count,
                r.status,r.correct_answers
                FROM math_results r JOIN math_tests t ON t.id=r.test_id
                WHERE r.student_id=%s ORDER BY t.held_on DESC,t.id DESC LIMIT 200""",
                (student_id,)).fetchall()
            return jsonify({"student":s,"history":history})

    @app.get("/api/math/revision")
    def math_revision_get():
        with db() as conn: return jsonify({"revision":revision(conn)})

    @app.get("/api/math/events")
    def math_events():
        # Enable only on a threaded server (gthread/gevent) with sufficient
        # request slots. Free/single-worker hosts use 6s revision polling.
        if os.environ.get("MATH_SSE_ENABLED") != "1":
            return bad("Real-time stream requires threaded VPS server",503)
        def generator():
            start=time.monotonic()
            last=-1
            while time.monotonic()-start<55:
                try:
                    with db() as conn:
                        v=revision(conn)
                    if v!=last:
                        yield "event: revision\ndata: " + str(v) + "\n\n"
                        last=v
                    else:
                        yield ": alive\n\n"
                except Exception:
                    yield "event: reconnect\ndata: retry\n\n"
                    return
                time.sleep(4)
        return Response(stream_with_context(generator()),mimetype="text/event-stream",
            headers={"Cache-Control":"no-cache, no-transform","X-Accel-Buffering":"no"})

    @app.post("/api/admin/math/commit")
    @admin_only
    def math_commit():
        obj=request.get_json(silent=True)
        if not isinstance(obj,dict): return bad("JSON ma’lumot kerak")
        date=date_value(obj.get("date"))
        title=text(obj,"title",120)
        try: questions=int(obj.get("question_count"))
        except (ValueError,TypeError): return bad("Savollar soni noto‘g‘ri")
        if not 1<=questions<=200:return bad("Savollar soni 1–200")
        entries=obj.get("entries")
        if not isinstance(entries,list) or not 1<=len(entries)<=MAX_ROWS:
            return bad("1–450 ta natija kiritilsin")
        parsed=[parse_row(x,questions) for x in entries]
        keys=[normalize(p["name"]) for p in parsed]
        if len(set(keys))!=len(keys):return bad("Bir xil ism bir martadan ko‘p kiritilgan")
        # Prevent a client from rewriting an existing test with a different
        # question count. This would make historical scores inconsistent.
        with db() as conn:
            with conn.transaction():
                old=conn.execute("""SELECT id,question_count FROM math_tests
                  WHERE held_on=%s AND title=%s FOR UPDATE""",(date,title)).fetchone()
                if old and old["question_count"]!=questions:
                    return bad("Bu sanada nomi bir xil test bor, savollar soni boshqa",409)
                test=(old or conn.execute("""INSERT INTO math_tests(held_on,title,question_count)
                   VALUES(%s,%s,%s) RETURNING id,question_count""",
                   (date,title,questions)).fetchone())
                for p in parsed:
                    name=p["name"]
                    st=conn.execute("""SELECT id,grade FROM math_students
                        WHERE lower(trim(full_name))=lower(trim(%s)) AND NOT archived
                        FOR UPDATE""",(name,)).fetchone()
                    if st is None:
                        st=conn.execute("""INSERT INTO math_students(full_name,grade)
                        VALUES(%s,%s) RETURNING id,grade""",(name,p["grade"])).fetchone()
                    elif p["grade"] is not None and st["grade"]!=p["grade"]:
                        conn.execute("""UPDATE math_students SET grade=%s WHERE id=%s""",
                                     (p["grade"],st["id"]))
                    conn.execute("""INSERT INTO math_results(test_id,student_id,status,correct_answers)
                      VALUES(%s,%s,%s,%s)
                      ON CONFLICT(test_id,student_id) DO UPDATE
                      SET status=EXCLUDED.status,correct_answers=EXCLUDED.correct_answers,
                          updated_at=now()""",
                      (test["id"],st["id"],p["status"],p["correct"]))
                conn.execute("""INSERT INTO math_import_audit(test_id,imported_rows,source)
                      VALUES(%s,%s,%s)""",(test["id"],len(parsed),
                       "bulk" if obj.get("source")=="bulk" else "manual"))
                bump(conn)
        return jsonify({"ok":True,"test_id":test["id"],"saved":len(parsed)}),201

    @app.post("/api/admin/math/payment")
    @admin_only
    def math_payment():
        obj=request.get_json(silent=True)
        if not isinstance(obj,dict):return bad("JSON ma’lumot kerak")
        try: student_id=int(obj.get("student_id"))
        except (TypeError,ValueError):return bad("O‘quvchi ID noto‘g‘ri")
        month=month_value(obj.get("month") or this_month().strftime("%Y-%m"))
        paid=obj.get("paid")
        if not isinstance(paid,bool):return bad("paid true yoki false bo‘lishi kerak")
        amount=obj.get("amount_uzs")
        if amount in ("",None):amount=None
        else:
            try: amount=int(amount)
            except (ValueError,TypeError):return bad("Summa noto‘g‘ri")
            if not 1000<=amount<=100000000:return bad("Summa 1000–100 mln oralig‘ida")
        note=str(obj.get("note") or "").strip()[:250]
        with db() as conn:
            with conn.transaction():
                student=conn.execute("""SELECT id FROM math_students
                    WHERE id=%s AND NOT archived""",(student_id,)).fetchone()
                if not student:return bad("O‘quvchi topilmadi",404)
                conn.execute("""INSERT INTO math_monthly_payments
                  (student_id,month,paid,amount_uzs,admin_note,paid_at)
                  VALUES(%s,%s,%s,%s,%s,CASE WHEN %s THEN now() ELSE NULL END)
                  ON CONFLICT(student_id,month) DO UPDATE SET
                  paid=EXCLUDED.paid,
                  amount_uzs=COALESCE(EXCLUDED.amount_uzs,math_monthly_payments.amount_uzs),
                  admin_note=EXCLUDED.admin_note,
                  paid_at=CASE WHEN EXCLUDED.paid THEN COALESCE(math_monthly_payments.paid_at,now()) ELSE NULL END,
                  updated_at=now()""",(student_id,month,paid,amount,note,paid))
                bump(conn)
        return jsonify({"ok":True,"paid":paid,"month":month.isoformat()})

    def clean_header(value):
        v=normalize(value).replace("‘","")
        return REQUIRED_HEADERS.get(v, None)

    def extract_rows(data, filename):
        ext=filename.rsplit(".",1)[-1].lower()
        if ext=="csv":
            try: content=data.decode("utf-8-sig")
            except UnicodeError: content=data.decode("cp1251")
            first=content[:4096]
            try: dialect=csv.Sniffer().sniff(first,delimiters=",;\t")
            except csv.Error: dialect=csv.excel
            return list(csv.reader(io.StringIO(content),dialect))[:MAX_ROWS+22]
        if ext=="xlsx":
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist())>220:raise ValueError("Excel ichida haddan tashqari ko‘p fayl")
                if sum(i.file_size for i in archive.infolist())>14_000_000:
                    raise ValueError("Excel fayli juda katta")
            wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
            try:
                if len(wb.worksheets)>12:raise ValueError("Excel varaqlari juda ko‘p")
                rows=[]
                for ws in wb.worksheets:
                    for row in ws.iter_rows(max_row=MAX_ROWS+20,max_col=18,values_only=True):
                        rows.append(list(row))
                        if len(rows)>=MAX_ROWS+20:break
                    if len(rows)>=MAX_ROWS+20:break
                return rows
            finally:wb.close()
        raise ValueError("Fayl .xlsx yoki .csv bo‘lishi kerak")

    def decode_preview(raw_rows):
        if not raw_rows:return [],["Jadval bo‘sh"]
        start=0
        columns={}
        for index,row in enumerate(raw_rows[:18]):
            matched={clean_header(val):pos for pos,val in enumerate(row)
                if clean_header(val) is not None}
            if "name" in matched and len(matched)>=2:
                columns=matched
                start=index+1
                break
        if not columns:
            columns={"name":0,"correct":1,"status":2,"grade":3}
        result=[]
        warnings=[]
        for rownum,row in enumerate(raw_rows[start:],start+1):
            if len(result)>=MAX_ROWS:break
            def col(key):
                idx=columns.get(key)
                return row[idx] if idx is not None and idx<len(row) else None
            name=col("name")
            if name is None or not str(name).strip():continue
            name=str(name).strip()
            if len(name)>NAME_LIMIT or len(name)<3:
                warnings.append("Qator "+str(rownum)+": ism noaniq")
                continue
            value=col("correct")
            state=str(col("status") or "").strip().lower()
            dq=(state in ("dq","d","absent","kelmadi","qatnashmadi") or
                str(value or "").strip().lower() in ("dq","d","kelmadi"))
            correct=None
            if not dq and value not in (None,""):
                try:
                    num=float(str(value).replace(",",".").replace("%",""))
                    if num.is_integer() and 0<=num<=200:correct=int(num)
                except (ValueError,TypeError):pass
            grade=col("grade")
            try:grade=int(grade) if grade not in (None,"") else None
            except (ValueError,TypeError):grade=None
            if grade is not None and not 2<=grade<=7:grade=None
            result.append({"name":name,"grade":grade,"correct":correct,
                "status":"dq" if dq else "present" if correct is not None else "unresolved",
                "row":rownum})
            if not dq and correct is None:
                warnings.append("Qator "+str(rownum)+": "+name+" uchun ball topilmadi")
        return result,warnings

    @app.post("/api/admin/math/preview")
    @admin_only
    def math_preview():
        incoming=request.files.getlist("files")
        if not 1<=len(incoming)<=8:return bad("1–8 ta fayl tanlang")
        out=[]
        for f in incoming:
            filename=(f.filename or "").rsplit("/",1)[-1][:150]
            ext=filename.rsplit(".",1)[-1].lower()
            data=f.read(2_100_000)
            if len(data)>2_000_000:
                out.append({"file":filename,"rows":[],"warnings":["Fayl 2 MB dan katta"]})
                continue
            if ext in ("jpg","jpeg","png","webp"):
                out.append({"file":filename,"rows":[],"warnings":[
                    "Rasm qabul qilindi, lekin avtomatik OCR servisi hali ulanmagan.",
                    "Buni natija deb saqlamaymiz. Excel/CSV yuboring yoki natijani qo‘lda to‘ldiring."]})
                continue
            try:
                records=extract_rows(data,filename)
                rows,warnings=decode_preview(records)
                out.append({"file":filename,"rows":rows,"warnings":warnings})
            except Exception as error:
                out.append({"file":filename,"rows":[],"warnings":[
                    "Faylni o‘qib bo‘lmadi: "+str(error)[:160]]})
        return jsonify({"files":out,"max_rows":MAX_ROWS,"requires_confirmation":True})
